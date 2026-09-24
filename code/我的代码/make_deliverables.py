from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from pptx import Presentation
from pptx.util import Inches as PInches, Pt as PPt
from pptx.dml.color import RGBColor as PRGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(r"D:\zijie")
RES = ROOT / "results_basic"
OUT = ROOT / "deliverables"
OUT.mkdir(exist_ok=True)

def load(name):
    return json.loads((RES / name).read_text(encoding="utf-8"))

base_m = load("baseline_metrics.json")
base_e = load("baseline_evaluation.json")
cmp_m = load("comparison_metrics.json")
cmp_e = load("comparison_evaluation.json")

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    from docx.oxml import OxmlElement
    shd = OxmlElement('w:shd')
    shd.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill', fill)
    tcPr.append(shd)

def add_doc_title(doc, text, subtitle=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text); r.bold = True; r.font.size = Pt(22); r.font.color.rgb = RGBColor(31,78,121)
    if subtitle:
        p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(subtitle); r2.font.size = Pt(11); r2.font.color.rgb = RGBColor(89,89,89)

def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.runs[0].font.color.rgb = RGBColor(31,78,121)
    return p

def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(item)

def add_image(doc, filename, width=6.1):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(RES / filename), width=Inches(width))

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.65); sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.75); sec.right_margin = Inches(0.75)
styles = doc.styles
styles['Normal'].font.name = 'Microsoft YaHei'; styles['Normal'].font.size = Pt(10.5)
add_doc_title(doc, '字节域语义内容理解：码流图像分类实验报告', 'ByteFormer 直接建模 JPEG 原始字节序列 · 基础实验与参数对比')
doc.add_paragraph('')
add_heading(doc, '摘要', 1)
doc.add_paragraph('本实验使用 ByteFormer Tiny 预训练模型，直接对 MNIST 图像 JPEG 文件的原始字节序列进行建模，不经过 RGB 解码，完成 0–9 十类数字分类。基础配置 batch size=32 时，验证集最佳准确率为 94.9%，独立测试集准确率为 94.4%（944/1000）。将 batch size 改为 16 后，测试准确率降至 93.3%，训练时间由 201.21 s 增至 218.84 s。')
add_heading(doc, '1. 实验背景与原理', 1)
doc.add_paragraph('传统视觉分类流程通常是“图像文件→解码为 RGB 像素→CNN/Transformer”。本实验把 JPEG 文件视为 0–255 范围内的字节序列，直接学习文件头、压缩数据和文件尾中的结构信息。JPEG 码流包含 SOI、DQT、SOF、DHT、SOS、熵编码图像数据以及 EOI 等部分；这些结构共同组成模型的输入上下文。')
doc.add_paragraph('ByteFormer 是面向原始字节序列的 Transformer 架构。本实验采用 ByteFormer Tiny 预训练权重，并将分类头调整为 10 类。训练阶段使用原图和轻微旋转、平移视图；验证集和测试集保持独立，不参与训练增强。')
add_heading(doc, '2. 数据集与实验设置', 1)
table = doc.add_table(rows=1, cols=3); table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.style = 'Table Grid'
for i, h in enumerate(['数据划分','数量','用途']):
    cell=table.rows[0].cells[i]; cell.text=h; set_cell_shading(cell,'1F4E79')
    for run in cell.paragraphs[0].runs: run.font.bold=True; run.font.color.rgb=RGBColor(255,255,255)
for row in [('训练集','5,000','更新模型参数'),('验证集','1,000','选择最佳模型'),('测试集','1,000','最终独立评估')]:
    cells=table.add_row().cells
    for i,v in enumerate(row): cells[i].text=v; cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
doc.add_paragraph('运行环境：Tesla T4 GPU，Python 3.12.13，PyTorch 2.10.0+cu128，随机种子 42，训练 10 轮，主干学习率 1e-4，weight decay 0.01，并在第 7、10 轮进行学习率衰减。')
add_heading(doc, '3. 实验流程', 1)
add_bullets(doc, [
    '准备数据、ByteFormer 预训练权重与运行环境。',
    '使用 batch size=32 进行基础训练，每轮在验证集上评估并保存最佳权重。',
    '使用独立测试集计算最终准确率、损失和混淆矩阵。',
    '将 batch size 改为 16，其他条件保持不变，进行参数对比。',
    '检查单图预测和典型错例，分析模型混淆来源。'
])
doc.add_paragraph('核心命令：')
for cmd in [
    'python train_course_subset.py --method clean --epochs 10 --batch-size 32 --clean-augmentations --output outputs/course_clean_v2',
    'python evaluate.py --checkpoint outputs/course_clean_v2/best.pt --output outputs/course_clean_v2_eval',
    'python predict.py --checkpoint outputs/course_clean_v2/best.pt --index 0']:
    p=doc.add_paragraph(); r=p.add_run(cmd); r.font.name='Consolas'; r.font.size=Pt(9)
add_heading(doc, '4. 基线实验结果', 1)
table = doc.add_table(rows=1, cols=2); table.style='Table Grid'; table.alignment=WD_TABLE_ALIGNMENT.CENTER
for i,h in enumerate(['指标','结果']):
    c=table.rows[0].cells[i]; c.text=h; set_cell_shading(c,'1F4E79')
    for run in c.paragraphs[0].runs: run.font.bold=True; run.font.color.rgb=RGBColor(255,255,255)
for k,v in [('最佳验证准确率','94.9%（第10轮）'),('最终训练准确率','97.58%'),('独立测试准确率','94.4%（944/1000）'),('测试损失','0.2225'),('训练耗时','201.21 s')]:
    cells=table.add_row().cells; cells[0].text=k; cells[1].text=v
add_image(doc, 'baseline_curves.png')
doc.add_paragraph('训练损失整体下降，训练准确率持续提升。第 8 轮学习率从 1e-4 降至 2e-5 后，验证准确率由 91.3% 提升至 94.3%，第 10 轮达到 94.9%。测试准确率与验证准确率接近，说明模型具有较好的泛化能力。')
add_heading(doc, '5. 参数对比实验', 1)
table=doc.add_table(rows=1, cols=5); table.style='Table Grid'; table.alignment=WD_TABLE_ALIGNMENT.CENTER
for i,h in enumerate(['实验','Epochs','Batch size','验证准确率','测试准确率']):
    c=table.rows[0].cells[i]; c.text=h; set_cell_shading(c,'1F4E79')
    for run in c.paragraphs[0].runs: run.font.bold=True; run.font.color.rgb=RGBColor(255,255,255)
for row in [('基线','10','32','94.9%','94.4%'),('对比','10','16','94.3%','93.3%')]:
    cells=table.add_row().cells
    for i,v in enumerate(row): cells[i].text=v
add_image(doc, 'comparison_curves.png')
doc.add_paragraph('batch size 从 32 减小到 16 后，验证准确率下降 0.6 个百分点，测试准确率下降 1.1 个百分点，训练时间增加约 17.63 s。当前设置下 batch size=32 收敛更快、最终性能更好，因此选择其作为最终基线。')
add_heading(doc, '6. 单图预测与错例分析', 1)
add_image(doc, 'prediction_single.png', width=4.8)
doc.add_paragraph('测试索引 0 的样本真实标签为 7，模型预测为 7，置信度约 99.96%。混淆矩阵显示模型共错分 56 个样本，主要混淆集中在 5/3、7/9 等笔画结构相似的数字之间。')
add_image(doc, 'baseline_confusion_matrix.png', width=5.5)
add_heading(doc, '7. 结论与后续工作', 1)
add_bullets(doc, [
    '直接对 JPEG 原始字节序列建模可以完成有效的 MNIST 十分类。',
    'ByteFormer Tiny 在 5,000 个训练样本上取得 94.4% 的独立测试准确率。',
    'batch size=32 优于 batch size=16，说明当前训练配置下较大的批量有利于收敛和泛化。',
    '后续可加入比特翻转和字节丢失增强，进一步评估模型对码流损坏的鲁棒性。'
])
add_heading(doc, '附录：提交文件清单', 1)
doc.add_paragraph('建议提交实验报告、基础模型 best.pt、训练指标与曲线、测试评估结果、参数对比结果以及单图预测图。')
doc_path = OUT / '字节域语义内容理解_实验报告.docx'
doc.save(doc_path)

# PPTX
prs = Presentation(); prs.slide_width=PInches(13.333); prs.slide_height=PInches(7.5)
BG=PRGBColor(247,249,252); NAVY=PRGBColor(31,78,121); BLUE=PRGBColor(46,117,182); DARK=PRGBColor(45,45,45); GRAY=PRGBColor(100,100,100); WHITE=PRGBColor(255,255,255); ORANGE=PRGBColor(237,125,49)

def add_bg(slide, title, kicker=None):
    bg=slide.background.fill; bg.solid(); bg.fore_color.rgb=BG
    bar=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,prs.slide_width,PInches(0.18)); bar.fill.solid(); bar.fill.fore_color.rgb=NAVY; bar.line.fill.background()
    tx=slide.shapes.add_textbox(PInches(0.55),PInches(0.42),PInches(12.1),PInches(0.6)); tf=tx.text_frame; tf.clear(); p=tf.paragraphs[0]; p.text=title; p.font.size=PPt(28); p.font.bold=True; p.font.color.rgb=NAVY; p.font.name='Microsoft YaHei'
    if kicker:
        k=slide.shapes.add_textbox(PInches(0.58),PInches(1.03),PInches(12),PInches(0.3)); p=k.text_frame.paragraphs[0]; p.text=kicker; p.font.size=PPt(11); p.font.color.rgb=GRAY; p.font.name='Microsoft YaHei'

def add_text(slide, x,y,w,h, text, size=18, color=DARK, bold=False, align=PP_ALIGN.LEFT):
    box=slide.shapes.add_textbox(PInches(x),PInches(y),PInches(w),PInches(h)); tf=box.text_frame; tf.word_wrap=True; tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.clear(); p=tf.paragraphs[0]; p.text=text; p.alignment=align; p.font.size=PPt(size); p.font.color.rgb=color; p.font.bold=bold; p.font.name='Microsoft YaHei'; return box

def add_card(slide,x,y,w,h,title,body,accent=BLUE):
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,PInches(x),PInches(y),PInches(w),PInches(h)); sh.fill.solid(); sh.fill.fore_color.rgb=WHITE; sh.line.color.rgb=PRGBColor(220,226,234)
    add_text(slide,x+0.22,y+0.18,w-0.4,0.35,title,16,NAVY,True); add_text(slide,x+0.22,y+0.7,w-0.4,h-0.85,body,13,DARK)

# 1 title
s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=NAVY
add_text(s,0.75,1.55,11.8,1.0,'字节域语义内容理解',34,WHITE,True,PP_ALIGN.CENTER)
add_text(s,0.75,2.62,11.8,0.7,'基于 ByteFormer 的码流图像分类实验',22,PRGBColor(220,235,250),False,PP_ALIGN.CENTER)
add_text(s,0.75,4.7,11.8,0.4,'实验报告与结果汇报  |  MNIST JPEG 原始字节序列',14,PRGBColor(200,215,235),False,PP_ALIGN.CENTER)

# 2 background
s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s,'1. 为什么直接理解字节流？','从“解码像素”转向“建模文件本身”')
add_card(s,0.7,1.55,3.7,3.6,'传统视觉流程','图像文件\n→ 解码为 RGB\n→ CNN / Vision Transformer\n→ 分类结果')
add_card(s,4.82,1.55,3.7,3.6,'本实验流程','JPEG 文件\n→ 0–255 字节序列\n→ ByteFormer\n→ 数字 0–9 分类',accent=ORANGE)
add_card(s,8.94,1.55,3.7,3.6,'核心意义','保留文件格式、压缩结构和字节上下文；不依赖预先解码的像素表示。')

# 3 structure
s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s,'2. 码流与模型输入','JPEG 结构信息共同构成字节域语义')
for x,t,b in [(0.8,'文件头','SOI / DQT / SOF / DHT / SOS\n记录尺寸、颜色分量和解码参数'),(4.75,'图像数据','DCT、量化和熵编码后的压缩数据\n以 MCU 等基本单元组织'),(8.7,'文件尾','EOI（FF D9）\n标志当前 JPEG 码流结束')]: add_card(s,x,1.7,3.65,3.1,t,b)
add_text(s,1.0,5.35,11.2,0.5,'ByteFormer 直接学习连续字节之间的上下文关系，分类头输出 10 个数字类别。',18,NAVY,True,PP_ALIGN.CENTER)

# 4 data
s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s,'3. 数据集与实验设置','固定划分，测试集只用于最终评估')
for x,num,label,desc in [(0.9,'5,000','训练集','更新模型参数'),(4.85,'1,000','验证集','选择最佳模型'),(8.8,'1,000','测试集','独立报告结果')]:
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,PInches(x),PInches(1.7),PInches(3.4),PInches(2.2)); sh.fill.solid(); sh.fill.fore_color.rgb=WHITE; sh.line.color.rgb=BLUE
    add_text(s,x+0.2,2.0,3.0,0.7,num,30,BLUE,True,PP_ALIGN.CENTER); add_text(s,x+0.2,2.8,3.0,0.4,label,17,NAVY,True,PP_ALIGN.CENTER); add_text(s,x+0.2,3.3,3.0,0.35,desc,13,GRAY,False,PP_ALIGN.CENTER)
add_text(s,1.1,4.75,11.0,1.0,'环境：Tesla T4 · Python 3.12.13 · PyTorch 2.10.0+cu128 · seed=42\n训练 10 轮 · lr=1e-4 · weight decay=0.01 · 第 7、10 轮衰减学习率',17,DARK,False,PP_ALIGN.CENTER)

# 5 process
s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s,'4. 实验流程','从准备资源到独立测试的完整闭环')
steps=[('01','准备','数据、依赖、预训练权重'),('02','训练','ByteFormer Tiny 微调'),('03','验证','每轮选择最佳权重'),('04','测试','独立测试集评估'),('05','对比','改变 batch size')]
for i,(n,t,b) in enumerate(steps):
    x=0.65+i*2.55; add_text(s,x,2.0,0.7,0.5,n,20,ORANGE,True,PP_ALIGN.CENTER); add_text(s,x-0.25,2.65,1.2,0.4,t,16,NAVY,True,PP_ALIGN.CENTER); add_text(s,x-0.55,3.2,1.8,0.9,b,12,DARK,False,PP_ALIGN.CENTER)
    if i<4: add_text(s,x+1.5,2.35,0.6,0.4,'→',24,GRAY,True,PP_ALIGN.CENTER)

# 6 baseline results
s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s,'5. 基线结果：batch size = 32','最佳验证准确率 94.9%，测试准确率 94.4%')
metrics=[('94.9%','最佳验证准确率'),('94.4%','测试准确率'),('944/1000','测试正确数'),('201.21 s','训练耗时')]
for i,(v,l) in enumerate(metrics): add_card(s,0.7+i*3.15,1.55,2.8,1.5,l,v)
s.shapes.add_picture(str(RES/'baseline_curves.png'),PInches(0.85),PInches(3.35),width=PInches(7.1))
add_text(s,8.3,3.55,4.0,2.3,'训练损失整体下降，准确率持续提升。\n\n第 8 轮学习率降至 2e-5 后，验证准确率进一步提升，最终在第 10 轮达到 94.9%。',16,DARK)

# 7 comparison
s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s,'6. 参数对比：batch size 的影响','其他条件保持一致，仅改变 batch size')
add_card(s,0.75,1.55,5.6,2.0,'基线：batch size = 32','验证：94.9%\n测试：94.4%\n训练时间：201.21 s')
add_card(s,6.95,1.55,5.6,2.0,'对比：batch size = 16','验证：94.3%\n测试：93.3%\n训练时间：218.84 s',accent=ORANGE)
s.shapes.add_picture(str(RES/'comparison_curves.png'),PInches(0.9),PInches(4.0),width=PInches(6.7))
add_text(s,8.0,4.2,4.1,1.5,'结论\n• batch=32 测试准确率高 1.1 个百分点\n• batch=16 训练时间反而增加\n• 最终选择 batch=32',16,NAVY,True)

# 8 prediction/error
s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s,'7. 单图预测与错误分析','正确样本置信度高，错误集中在形态相似数字')
s.shapes.add_picture(str(RES/'prediction_single.png'),PInches(0.85),PInches(1.55),width=PInches(4.2))
s.shapes.add_picture(str(RES/'baseline_confusion_matrix.png'),PInches(5.35),PInches(1.55),width=PInches(4.7))
add_text(s,10.25,1.9,2.3,2.8,'样例：\n真实标签 7\n预测标签 7\n置信度约 99.96%\n\n主要混淆：\n5↔3、7↔9',16,DARK)

# 9 conclusion
s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s,'8. 结论与后续工作','基础实验已完成，结果满足课程要求')
add_bullets_ppt=[]
for i,t in enumerate(['JPEG 原始字节序列可以直接用于 MNIST 分类。','ByteFormer Tiny 在独立测试集达到 94.4% 准确率。','batch size=32 在本设置下优于 batch size=16。','后续可加入比特翻转、字节丢失增强，研究码流鲁棒性。']):
    add_text(s,1.0,1.65+i*0.9,11.0,0.5,'• '+t,20,NAVY if i<3 else ORANGE,True if i==1 else False)

# 10 deliverables
s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s,'9. 交付物','报告、模型、曲线与评估结果')
add_card(s,0.8,1.55,5.6,3.5,'实验报告','实验背景\n数据与环境\n训练流程\n结果与参数对比\n错例分析与结论')
add_card(s,6.95,1.55,5.6,3.5,'结果文件','best.pt\nmetrics.json / history.csv\ncurves.png\nevaluation.json\nconfusion_matrix.png\nprediction_single.png')
add_text(s,1.0,5.7,11.2,0.4,'谢谢！',26,NAVY,True,PP_ALIGN.CENTER)

ppt_path=OUT/'字节域语义内容理解_实验汇报.pptx'
prs.save(ppt_path)
print(doc_path)
print(ppt_path)
